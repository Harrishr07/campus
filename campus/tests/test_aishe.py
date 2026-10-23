"""M01-09 AISHE dataset export: sections built from verified records equal their source data points and reports,
the layout is a placeholder mapping record, cells are formula-safe, and the export stays inside the user's company."""

import io

import frappe
from openpyxl import load_workbook

from campus.campus_finance.report.scholarship_beneficiaries import scholarship_beneficiaries as beneficiaries_report
from campus.examinations.report.result_analysis import result_analysis
from campus.iqac import aishe, registry
from campus.iqac.official_seed import seed_official_register
from campus.tests import demo
from campus.tests.demo_faculty import YEAR
from campus.tests import test_scholarships as scholarship_year
from campus.tests.factories import make_user
from campus.tests.test_iqac import COLLEGE, IqacFixture
from campus.tests.test_regulation import year_name
from campus.tests.test_scholarships import COLLEGE as SCHOLARSHIP_COLLEGE, ScholarshipFixture
from campus.tests.utils import CampusTestCase

INJECTION = '=HYPERLINK("http://evil.example","x")'
COORDINATOR, OUTSIDER, PLAIN = "aishe-coordinator@example.com", "aishe-outsider@example.com", "aishe-plain@example.com"
QUOTA = "AISHE Test Quota"


def point(key):
	return registry.data_point(key, demo.ENGINEERING, YEAR)


class TestLayoutRecord(CampusTestCase):
	def setUp(self):
		frappe.db.savepoint("aishe_layout")
		self.addCleanup(frappe.db.rollback, save_point="aishe_layout")
		seed_official_register()

	def layout(self, sections):
		return frappe.get_doc(
			{"doctype": "AISHE Layout", "layout_name": f"Test Layout {frappe.generate_hash(length=5)}", "sections": sections}
		)

	def test_the_seeded_layout_is_labelled_a_placeholder_and_maps_every_section(self):
		layout = frappe.get_doc("AISHE Layout", aishe.PLACEHOLDER_LAYOUT)
		self.assertTrue(layout.is_placeholder)
		self.assertEqual({row.section for row in layout.sections}, set(aishe.SECTIONS))

	def test_the_section_choices_are_the_sections_the_export_builds(self):
		options = frappe.get_meta("AISHE Layout Section").get_field("section").options.split("\n")
		self.assertEqual(options, aishe.SECTIONS)

	def test_a_bad_sheet_name_or_cell_is_refused(self):
		for row in ({"section": "Institution", "sheet": "A/B", "first_cell": "A1"}, {"section": "Institution", "sheet": "Basic", "first_cell": "ZZZZ9"}):
			with self.subTest(row=row), self.assertRaises(frappe.ValidationError):
				self.layout([row]).insert()

	def test_a_section_is_mapped_once_and_two_sections_do_not_start_in_the_same_cell(self):
		same_section = [{"section": "Institution", "sheet": "A", "first_cell": "A1"}, {"section": "Institution", "sheet": "B", "first_cell": "A1"}]
		same_cell = [{"section": "Institution", "sheet": "A", "first_cell": "a1"}, {"section": "Faculty", "sheet": "A", "first_cell": "A1"}]
		for rows in (same_section, same_cell):
			with self.subTest(rows=rows), self.assertRaises(frappe.ValidationError):
				self.layout(rows).insert()


class AisheCase(CampusTestCase):
	demo_scale = "small"

	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		for user, roles, companies in ((COORDINATOR, ["IQAC Coordinator"], [demo.ENGINEERING]), (OUTSIDER, ["IQAC Coordinator"], [demo.ARTS]), (PLAIN, [], [demo.ENGINEERING])):
			make_user(user, companies)
			frappe.get_doc("User", user).add_roles(*roles)

	def setUp(self):
		frappe.db.savepoint("aishe")
		self.addCleanup(frappe.db.rollback, save_point="aishe")
		self.addCleanup(frappe.set_user, "Administrator")
		seed_official_register()
		self.enrolments = frappe.get_all(
			"Program Enrollment", {"oc_company": demo.ENGINEERING, "academic_year": YEAR, "docstatus": 1}, ["student", "program", "student_category"], order_by="name"
		)
		for number, row in enumerate(self.enrolments):  # gender and state alternate so the split has several cells
			frappe.db.set_value("Student", row.student, {"gender": ("Male", "Female")[number % 2], "state": ("Tamil Nadu", "Kerala", "")[number % 3]})

	def rows(self, section):
		return aishe.build(demo.ENGINEERING, YEAR)[section]["rows"]

	def workbook(self, layout=None):
		_name, content = aishe.export_dataset(demo.ENGINEERING, YEAR, layout)
		return load_workbook(io.BytesIO(content))


class TestSections(AisheCase):
	def test_enrolment_splits_distinct_students_by_category_gender_and_state(self):
		rows = self.rows("Enrolment")
		self.assertEqual(sum(r[3] for r in rows), point("total_students")["value"])
		students = {r.student: r for r in self.enrolments}
		hand = {}
		for number, row in enumerate(self.enrolments):
			key = (row.student_category, ("Male", "Female")[number % 2], ("Tamil Nadu", "Kerala", aishe.NOT_RECORDED)[number % 3])
			hand[key] = hand.get(key, 0) + 1
		self.assertEqual({(r[0], r[1], r[2]): r[3] for r in rows}, hand)
		self.assertEqual(len(students), sum(hand.values()))

	def test_programmes_list_intake_and_admitted_per_program(self):
		if not frappe.db.exists("Admission Quota", QUOTA):
			frappe.get_doc({"doctype": "Admission Quota", "code": QUOTA, "title": QUOTA}).insert()
		frappe.get_doc(
			{"doctype": "Seat Matrix", "company": demo.ENGINEERING, "program": "B.E. CSE", "academic_year": YEAR, "rows": [{"quota": QUOTA, "category": "OC", "sanctioned": 12}, {"quota": QUOTA, "category": "BC", "sanctioned": 8}]}
		).insert()
		rows = {r[0]: r for r in self.rows("Programmes")}
		self.assertEqual(rows["B.E. CSE"][2], 20)  # 12 + 8 sanctioned
		self.assertEqual(sum(r[3] for r in rows.values()), point("students_admitted")["value"])
		self.assertEqual(len([r for r in rows.values() if r[2]]), point("programmes_offered")["value"])
		self.assertEqual(rows["B.E. CSE"][1], frappe.db.get_value("Program", "B.E. CSE", "department"))

	def test_faculty_splits_the_teachers_on_roll_by_designation_employment_type_and_gender(self):
		rows = self.rows("Faculty")
		self.assertEqual(sum(r[3] for r in rows), point("full_time_teachers")["value"])
		self.assertEqual(sum(r[3] for r in rows), 4)

	def test_institution_rows_carry_the_aishe_code_of_the_profile(self):
		frappe.db.set_value("Institution Profile", {"company": demo.ENGINEERING}, "aishe_code", "C-12345")
		rows = dict(self.rows("Institution"))
		self.assertEqual(rows["AISHE Code"], "C-12345")
		self.assertEqual(rows["Institution"], demo.ENGINEERING)

	def test_an_empty_source_gives_a_header_and_no_rows(self):
		self.assertEqual(aishe.build(demo.ENGINEERING, YEAR)["Examination Results"]["rows"], [])


class TestWorkbook(AisheCase):
	def test_the_cover_labels_the_layout_a_placeholder_and_names_the_aishe_code(self):
		book = self.workbook()
		cover = [[c.value for c in row] for row in book[aishe.COVER].iter_rows()]
		self.assertTrue(any("Placeholder" in str(cell) for row in cover for cell in row))
		self.assertEqual(book.sheetnames[0], aishe.COVER)

	def test_every_mapped_section_is_written_where_the_layout_says(self):
		layout = frappe.get_doc("AISHE Layout", aishe.PLACEHOLDER_LAYOUT)
		row = next(r for r in layout.sections if r.section == "Enrolment")
		book = self.workbook()
		sheet = book[row.sheet]
		first = sheet[row.first_cell]
		self.assertEqual(first.value, aishe.build(demo.ENGINEERING, YEAR)["Enrolment"]["header"][0])
		below = sheet.cell(row=first.row + 1, column=first.column).value
		self.assertEqual(below, self.rows("Enrolment")[0][0])

	def test_text_that_would_run_as_a_formula_is_written_as_text(self):
		frappe.db.set_value("Student", self.enrolments[0].student, "state", "+cmd|' /C calc'!A0")
		frappe.db.set_value("Student", self.enrolments[1].student, "state", INJECTION)
		book = self.workbook()
		cells = [c for sheet in book for row in sheet.iter_rows() for c in row if c.value is not None]
		self.assertTrue(cells)
		for cell in cells:
			self.assertNotEqual(cell.data_type, "f", cell.coordinate)
			if isinstance(cell.value, str):
				self.assertNotIn(cell.value.lstrip(" ")[:1], ("=", "+", "-", "@", "\t", "\r"), cell.coordinate)
		self.assertTrue(any(isinstance(c.value, str) and c.value.startswith("'+cmd") for c in cells))
		self.assertTrue(any(isinstance(c.value, str) and c.value.startswith("'=HYPERLINK") for c in cells))


class TestOverlap(AisheCase):
	def test_two_sections_whose_tables_overlap_are_refused_at_export(self):
		layout = frappe.get_doc(
			{"doctype": "AISHE Layout", "layout_name": "Overlap Layout", "sections": [
				{"section": "Institution", "sheet": "One", "first_cell": "A1"}, {"section": "Programmes", "sheet": "One", "first_cell": "A2"}]}
		).insert()
		with self.assertRaises(frappe.ValidationError):
			self.workbook(layout.name)


class TestScope(AisheCase):
	def test_the_export_is_for_iqac_staff_of_the_company_only(self):
		frappe.set_user(COORDINATOR)
		_name, content = aishe.export_dataset(demo.ENGINEERING, YEAR)  # a coordinator without HR User still gets every teacher
		sheet = load_workbook(io.BytesIO(content))["Faculty"]
		self.assertEqual(sum(row[3].value for row in sheet.iter_rows(min_row=4) if isinstance(row[3].value, int)), 4)
		frappe.set_user(OUTSIDER)  # an IQAC Coordinator of another college
		with self.assertRaises(frappe.PermissionError):
			aishe.export_dataset(demo.ENGINEERING, YEAR)
		frappe.set_user(PLAIN)  # staff of the college without an IQAC role
		with self.assertRaises(frappe.PermissionError):
			aishe.export_dataset(demo.ENGINEERING, YEAR)

	def test_another_colleges_records_are_not_in_the_numbers(self):
		other = aishe.build(demo.ARTS, YEAR)
		self.assertEqual(sum(r[3] for r in other["Enrolment"]["rows"]), registry.data_point("total_students", demo.ARTS, YEAR)["value"])
		self.assertEqual(sum(r[3] for r in self.rows("Enrolment")), point("total_students")["value"])


class TestExaminationResults(IqacFixture):
	def test_final_year_results_equal_the_result_analysis_report_and_the_data_points(self):
		self.make_results()
		year = year_name("2022")
		filters = frappe._dict(company=COLLEGE, academic_year=year, final_year_only=1, exam_session=None)
		report = [(r["group"], r["appeared"], r["passed"], r["pass_pct"]) for r in result_analysis.data(filters, "Program")]
		section = aishe.build(COLLEGE, year)["Examination Results"]
		self.assertEqual([tuple(r) for r in section["rows"]], report)
		total = section["rows"][-1]
		self.assertEqual((total[1], total[2]), (3, 2))  # hand count of the fixture: s22, s26, s_ab appeared; s22, s26 passed
		self.assertEqual(total[1], registry.data_point("final_year_appeared", COLLEGE, year)["value"])
		self.assertEqual(total[2], registry.data_point("final_year_passed", COLLEGE, year)["value"])
		self.assertEqual(total[3], registry.data_point("pass_percentage", COLLEGE, year)["value"])

	def test_unpublished_results_are_not_in_the_dataset(self):
		self.make_results()
		frappe.db.set_value("Exam Session", self.session.name, "status", "Marks Entry")
		self.assertEqual(aishe.build(COLLEGE, year_name("2022"))["Examination Results"]["rows"], [])


class TestScholarships(ScholarshipFixture):
	def test_students_per_scheme_type_equal_the_data_points_and_the_report(self):
		self.application(self.a, self.govt)
		self.application(self.b, self.govt)
		self.application(self.b, self.inst)
		self.application(self.c, self.inst, submit=False)  # a draft is not a sanction
		frappe.db.set_value("Student", self.a, {"gender": "Male"})
		rows = aishe.build(SCHOLARSHIP_COLLEGE, scholarship_year.YEAR)["Scholarships"]["rows"]
		by_type = {kind: sum(r[3] for r in rows if r[0] == kind) for kind in ("Government", "Institutional")}
		self.assertEqual(by_type, {"Government": 2, "Institutional": 1})
		self.assertEqual(by_type["Government"], registry.data_point("scholarship_govt_students", SCHOLARSHIP_COLLEGE, scholarship_year.YEAR)["value"])
		self.assertEqual(by_type["Institutional"], registry.data_point("scholarship_inst_students", SCHOLARSHIP_COLLEGE, scholarship_year.YEAR)["value"])
		summary = {s["label"]: s["value"] for s in beneficiaries_report.execute({"company": SCHOLARSHIP_COLLEGE, "academic_year": scholarship_year.YEAR})[4]}
		self.assertEqual(sorted(summary.values()), [1, 2])
