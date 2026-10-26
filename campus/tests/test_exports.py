"""M01-03 (P3-T06) Export Templates (nodes -> sheet/cell or section as data) and the AQAR Excel / PDF exporter."""

import io
import json

import frappe
from frappe.tests import UnitTestCase
from openpyxl import load_workbook
from pypdf import PdfReader

from campus.iqac import api, compute, export
from campus.iqac import cycle as lifecycle
from campus.iqac.export_seed import AQAR_EXCEL, AQAR_PDF, NBA_EXCEL, NIRF_EXCEL, seed_export_templates
from campus.iqac.setup import FRAMEWORK
from campus.tests.test_iqac import COORDINATOR, FACULTY_USER, OTHER_COLLEGE, OTHER_COORDINATOR, IqacFixture
from campus.tests.test_results import COLLEGE

N263, N211, N212, N111, N141 = (
	f"{FRAMEWORK}-{code}" for code in ("2.6.3", "2.1.1", "2.1.2", "1.1.1", "1.4.1")
)
EP21 = f"{FRAMEWORK}-EP-2.1"
CHOICE = "A. All 4 of the above"
EXPORTER = "iqac-exporter@example.com"
INJECTION = '=HYPERLINK("http://evil.example","x")'


class TestCellSafety(UnitTestCase):
	def test_text_that_would_run_as_a_formula_is_neutralised(self):
		for text in ("=1+1", "+1", "-1", "@SUM(A1)", "\t=1", "\r=1", " =1+1", "  @x"):
			self.assertEqual(export.safe_text(text), "'" + text)

	def test_plain_text_is_unchanged_and_none_is_blank(self):
		self.assertEqual(export.safe_text("Students admitted"), "Students admitted")
		self.assertEqual(export.safe_text(None), "")
		self.assertEqual(export.safe_text("a=b"), "a=b")

	def test_control_characters_are_dropped_and_length_is_capped(self):
		self.assertEqual(export.safe_text("a\x00b\x0bc\nd"), "abc\nd")
		self.assertEqual(len(export.safe_text("x" * 40000)), export.MAX_CELL)

	def test_sheet_names_lose_forbidden_characters_and_fit_31(self):
		self.assertEqual(export.sheet_name("A/B:C*D?E[F]G\\H"), "A B C D E F G H")
		self.assertEqual(len(export.sheet_name("x" * 50)), 31)
		self.assertEqual(export.sheet_name("'quoted'"), "quoted")
		self.assertEqual(export.sheet_name("///"), "Sheet")

	def test_narrative_is_plain_text(self):
		self.assertEqual(export.plain("<p>One &amp; <b>two</b></p><p>three</p>"), "One & two three")


class ExportFixture(IqacFixture):
	def setUp(self):
		super().setUp()
		seed_export_templates()

	def lock(self, cycle=None):
		"""Compute, answer a few nodes by hand, lock the values and approve the cycle (direct writes: this is test setup)."""
		cycle = cycle or self.cycle
		self.make_results()
		compute.compute_cycle(cycle.name)
		for node, values in (
			(N212, {"manual_value": 12.5}), (EP21, {"manual_value": 321}), (N141, {"choice": CHOICE}),
			(N111, {"narrative": "<p>Plan one</p><p>Second</p>"}),
		):  # fmt: skip
			frappe.get_doc({"doctype": "Metric Value", "cycle": cycle.name, "node": node, **values}).insert()
		lifecycle.lock_values(cycle.name)
		frappe.db.set_value("Accreditation Cycle", cycle.name, "status", "Approved")

	def workbook(self, template=AQAR_EXCEL, cycle=None):
		_, content = export.export((cycle or self.cycle).name, template)
		return load_workbook(io.BytesIO(content))

	def mapped_cell(self, book, template, node):
		row = next(r for r in frappe.get_doc("Export Template", template).mappings if r.node == node)
		return book[row.sheet][row.cell].value

	def custom_template(self, rows, **values):
		return frappe.get_doc(
			{
				"doctype": "Export Template", "template_name": f"Test Template {frappe.generate_hash(length=6)}",
				"framework": FRAMEWORK, "export_format": "Excel", "mappings": rows, **values,
			}
		).insert()  # fmt: skip


class TestExcelValues(ExportFixture):
	def test_exported_values_equal_the_locked_metric_values(self):
		self.lock()
		book = self.workbook()
		checked = 0
		for value in frappe.get_all(
			"Metric Value",
			{"cycle": self.cycle.name, "department": ["is", "not set"]},
			["node", "value", "choice", "status", "answered"],
		):
			self.assertEqual(value.status, "Locked")
			answer_type = frappe.db.get_value("Framework Node", value.node, "answer_type")
			if answer_type in ("Number", "Percent"):
				self.assertEqual(self.mapped_cell(book, AQAR_EXCEL, value.node), value.value, value.node)
				checked += 1
		self.assertEqual(self.mapped_cell(book, AQAR_EXCEL, N263), 66.67)
		self.assertEqual(self.mapped_cell(book, AQAR_EXCEL, N211), 4)
		self.assertEqual(self.mapped_cell(book, AQAR_EXCEL, N212), 12.5)
		self.assertEqual(self.mapped_cell(book, AQAR_EXCEL, EP21), 321)
		self.assertEqual(self.mapped_cell(book, AQAR_EXCEL, N141), CHOICE)
		self.assertGreaterEqual(checked, 4)

	def test_a_narrative_goes_out_as_plain_text_and_unanswered_nodes_stay_blank(self):
		self.lock()
		book = self.workbook()
		self.assertIn("Plan one", self.mapped_cell(book, AQAR_EXCEL, N111))
		self.assertNotIn("<p>", self.mapped_cell(book, AQAR_EXCEL, N111))
		unanswered = f"{FRAMEWORK}-2.1.2.1"
		if any(r.node == unanswered for r in frappe.get_doc("Export Template", AQAR_EXCEL).mappings):
			self.assertIsNone(self.mapped_cell(book, AQAR_EXCEL, unanswered))

	def test_the_data_template_sheet_lists_the_snapshot_records(self):
		self.lock()
		book = self.workbook()
		sheet = book[export.data_sheet_name("2.6.3")]
		rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=1)]
		header = next(i for i, row in enumerate(rows) if row[:1] == ["Department"])
		listed = {(row[1], row[4]) for row in rows[header + 1 :] if row[4]}
		snapshot = json.loads(frappe.db.get_value("Metric Value", f"{self.cycle.name}-2.6.3", "snapshot"))
		expected = {
			(key, record["name"]) for key, point in snapshot["inputs"].items() for record in point["records"]
		}
		self.assertTrue(expected)
		self.assertTrue(expected <= listed, expected - listed)
		self.assertEqual({k for k, _ in expected}, {"final_year_passed", "final_year_appeared"})

	def test_records_the_exporter_cannot_read_are_counted_not_listed(self):
		self.lock()
		snapshot = json.loads(frappe.db.get_value("Metric Value", f"{self.cycle.name}-2.6.3", "snapshot"))
		results = {
			r["name"]
			for point in snapshot["inputs"].values()
			for r in point["records"]
			if r["doctype"] == "Student Result"
		}
		self.assertTrue(results)
		if not frappe.db.exists(
			"User", EXPORTER
		):  # IQAC Coordinator alone: no exam role, so no Student Result read
			frappe.get_doc(
				{
					"doctype": "User",
					"email": EXPORTER,
					"first_name": "exporter",
					"send_welcome_email": 0,
					"roles": [{"role": "IQAC Coordinator"}],
				}
			).insert()
			frappe.get_doc(
				{"doctype": "User Permission", "user": EXPORTER, "allow": "Company", "for_value": COLLEGE}
			).insert()
		frappe.set_user(EXPORTER)
		self.assertFalse(frappe.has_permission("Student Result", "read"))
		book = self.workbook()
		cells = [
			c.value
			for row in book[export.data_sheet_name("2.6.3")].iter_rows()
			for c in row
			if c.value is not None
		]
		self.assertFalse(results & set(cells), "unreadable Student Results are listed")
		self.assertIn(f"{len(results)} records you cannot open are not listed.", cells)

	def test_a_snapshot_point_without_records_is_listed_once_without_a_record(self):
		self.lock()
		name = f"{self.cycle.name}-2.6.3"
		snapshot = json.loads(frappe.db.get_value("Metric Value", name, "snapshot"))
		snapshot["inputs"]["bare_point"] = {"value": 3}  # no "records" key
		frappe.db.set_value("Metric Value", name, "snapshot", json.dumps(snapshot))
		rows = [
			[c.value for c in row] for row in self.workbook()[export.data_sheet_name("2.6.3")].iter_rows()
		]
		self.assertEqual([row[4] for row in rows if row[1] == "bare_point"], [None])

	def test_the_cover_sheet_says_what_was_exported(self):
		self.lock()
		cover = self.workbook()[export.COVER]
		text = " ".join(str(c.value) for row in cover.iter_rows() for c in row if c.value)
		self.assertIn(self.cycle.name, text)
		self.assertIn("placeholder", text.lower())
		self.assertNotIn("not Locked", text)

	def test_an_unlocked_cycle_is_marked_on_the_cover(self):
		compute.compute_cycle(self.cycle.name)
		text = " ".join(
			str(c.value) for row in self.workbook()[export.COVER].iter_rows() for c in row if c.value
		)
		self.assertIn("not Locked", text)


class TestMappingIsData(ExportFixture):
	def test_moving_a_cell_moves_the_value_without_code_changes(self):
		self.lock()
		template = self.custom_template([{"node": N263, "sheet": "Results", "cell": "B2"}])
		book = self.workbook(template.name)
		self.assertEqual(book["Results"]["B2"].value, 66.67)
		template.mappings[0].cell = "D7"
		template.mappings[0].sheet = "Moved"
		template.save()
		book = self.workbook(template.name)
		self.assertEqual(book["Moved"]["D7"].value, 66.67)
		self.assertNotIn("Results", book.sheetnames)

	def test_a_label_cell_gets_the_node_code_and_title(self):
		self.lock()
		template = self.custom_template([{"node": N263, "sheet": "R", "cell": "B2", "label_cell": "A2"}])
		self.assertTrue(self.workbook(template.name)["R"]["A2"].value.startswith("2.6.3"))

	def test_data_sheet_flag_decides_which_nodes_get_a_listing(self):
		self.lock()
		template = self.custom_template(
			[
				{"node": N263, "sheet": "R", "cell": "B2"},
				{"node": N211, "sheet": "R", "cell": "B3", "data_sheet": 1},
			]
		)
		names = self.workbook(template.name).sheetnames
		self.assertIn(export.data_sheet_name("2.1.1"), names)
		self.assertNotIn(export.data_sheet_name("2.6.3"), names)

	def test_bad_mappings_are_refused(self):
		for row in (
			{"node": N263, "sheet": "R", "cell": "A0"}, {"node": N263, "sheet": "R", "cell": "ZZZZ1"},
			{"node": N263, "sheet": "R", "cell": "$B2"}, {"node": N263, "sheet": "R"}, {"node": N263, "cell": "B2"},
			{"node": N263, "sheet": "Cover", "cell": "B2"}, {"node": N263, "sheet": "Data 2.6.3", "cell": "B2"},
			{"node": N263, "sheet": "a/b", "cell": "B2"}, {"node": N263, "sheet": "x" * 32, "cell": "B2"},
			{"node": N263, "sheet": "R", "cell": "B2", "label_cell": "B2"},
		):  # fmt: skip
			with self.assertRaises(frappe.ValidationError, msg=str(row)):
				self.custom_template([row])

	def test_two_nodes_cannot_share_a_cell(self):
		with self.assertRaises(frappe.ValidationError):
			self.custom_template(
				[{"node": N263, "sheet": "R", "cell": "B2"}, {"node": N211, "sheet": "r", "cell": "b2"}]
			)

	def test_a_node_of_another_framework_is_refused(self):
		other = frappe.get_doc(
			{"doctype": "Accreditation Framework", "framework_name": "Export Other FW"}
		).insert()
		node = frappe.get_doc(
			{
				"doctype": "Framework Node",
				"framework": other.name,
				"code": "9",
				"title": "x",
				"node_type": "Criterion",
			}
		).insert()
		with self.assertRaises(frappe.ValidationError):
			self.custom_template([{"node": node.name, "sheet": "R", "cell": "B2"}])

	def test_a_pdf_template_needs_a_section_and_an_excel_one_a_cell(self):
		with self.assertRaises(frappe.ValidationError):
			self.custom_template([{"node": N263, "sheet": "R", "cell": "B2"}], export_format="PDF")
		self.assertTrue(self.custom_template([{"node": N263, "section": "Part B"}], export_format="PDF").name)


class TestFormulaInjection(ExportFixture):
	def test_no_cell_of_any_sheet_is_a_formula_and_user_text_is_neutralised(self):
		self.lock()
		frappe.db.set_value("Metric Value", f"{self.cycle.name}-1.1.1", "narrative", f"<p>{INJECTION}</p>")
		name = f"{self.cycle.name}-2.6.3"
		snapshot = json.loads(frappe.db.get_value("Metric Value", name, "snapshot"))
		snapshot["inputs"][INJECTION] = {
			"value": 1,
			"records": [],
		}  # a snapshot key, written to the Data Template sheet
		frappe.db.set_value("Metric Value", name, "snapshot", json.dumps(snapshot))
		book = self.workbook()
		texts = []
		for sheet in book.worksheets:
			for row in sheet.iter_rows():
				for cell in row:
					self.assertNotEqual(cell.data_type, "f", f"{sheet.title}!{cell.coordinate}")
					if isinstance(cell.value, str):
						texts.append(cell.value)
						self.assertNotIn(cell.value.lstrip(" ")[:1], ("=", "+", "@", "\t", "\r"), cell.value)
		self.assertIn("'" + INJECTION, texts)  # the narrative and the data point key both appear, defused
		self.assertGreaterEqual(texts.count("'" + INJECTION), 2)

	def test_a_formula_like_sheet_name_or_label_cannot_break_out(self):
		self.lock()
		template = self.custom_template([{"node": N263, "sheet": "cmd", "cell": "B2"}])
		frappe.db.set_value(
			"Export Template Mapping", template.mappings[0].name, "sheet", "=cmd"
		)  # past the controller's checks
		book = self.workbook(template.name)
		self.assertTrue(all(not name.startswith("=") for name in book.sheetnames))


class TestPermissionsAndScope(ExportFixture):
	def test_only_iqac_staff_and_the_principal_may_export(self):
		self.lock()
		frappe.set_user(FACULTY_USER)
		with self.assertRaises(frappe.PermissionError):
			api.export_cycle(self.cycle.name, AQAR_EXCEL)
		frappe.set_user(COORDINATOR)
		api.export_cycle(self.cycle.name, AQAR_EXCEL)
		self.assertEqual(frappe.response["type"], "download")
		self.assertTrue(frappe.response["filename"].endswith(".xlsx"))
		self.assertTrue(frappe.response["filecontent"].startswith(b"PK"))

	def test_a_coordinator_cannot_export_another_colleges_cycle(self):
		other = self.make_cycle(OTHER_COLLEGE)
		frappe.set_user(COORDINATOR)
		with self.assertRaises(frappe.PermissionError):
			api.export_cycle(other.name, AQAR_EXCEL)
		frappe.set_user(OTHER_COORDINATOR)
		api.export_cycle(other.name, AQAR_EXCEL)

	def test_the_template_must_belong_to_the_cycles_framework(self):
		self.lock()
		with self.assertRaises(frappe.ValidationError):
			export.export(self.cycle.name, NBA_EXCEL)

	def test_an_export_is_noted_on_the_cycle(self):
		self.lock()
		frappe.set_user(COORDINATOR)
		api.export_cycle(self.cycle.name, AQAR_EXCEL)
		comments = frappe.get_all(
			"Comment",
			{
				"reference_doctype": "Accreditation Cycle",
				"reference_name": self.cycle.name,
				"content": ["like", f"%{AQAR_EXCEL}%"],
			},
		)
		self.assertEqual(len(comments), 1)


class TestPdf(ExportFixture):
	def text(self, cycle=None):
		_, content = export.export((cycle or self.cycle).name, AQAR_PDF)
		self.assertTrue(content.startswith(b"%PDF"))
		return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)

	def test_the_pdf_follows_part_a_extended_profile_part_b_and_shows_the_locked_values(self):
		self.lock()
		text = self.text()
		part_a, extended, part_b = text.index("Part A"), text.index("Extended Profile"), text.index("Part B")
		self.assertLess(part_a, extended)
		self.assertLess(extended, part_b)
		self.assertIn("66.67", text)
		self.assertIn(CHOICE, text)
		self.assertIn("Plan one", text)
		self.assertIn(self.cycle.name, text)

	def test_markup_in_answers_is_not_rendered_into_the_pdf(self):
		self.lock()
		frappe.db.set_value(
			"Metric Value",
			f"{self.cycle.name}-1.1.1",
			"narrative",
			"<p>Hello <script>alert(1)</script><b>bold</b></p>",
		)
		text = self.text()
		self.assertNotIn("<b>", text)
		self.assertIn("Hello", text)


class TestSeeds(ExportFixture):
	def test_the_regulator_templates_are_seeded_as_labelled_placeholders(self):
		for name, fmt in (
			(AQAR_EXCEL, "Excel"),
			(AQAR_PDF, "PDF"),
			(NBA_EXCEL, "Excel"),
			(NIRF_EXCEL, "Excel"),
		):
			template = frappe.get_doc("Export Template", name)
			self.assertEqual(template.export_format, fmt)
			self.assertTrue(template.is_placeholder)
			self.assertIn("placeholder", (template.description or "").lower())
		self.assertGreater(len(frappe.get_doc("Export Template", AQAR_EXCEL).mappings), 50)

	def test_seeding_again_keeps_the_colleges_edits(self):
		template = frappe.get_doc("Export Template", AQAR_EXCEL)
		template.mappings[0].cell = "H99"
		template.save()
		seed_export_templates()
		self.assertEqual(frappe.get_doc("Export Template", AQAR_EXCEL).mappings[0].cell, "H99")
		self.assertEqual(frappe.db.count("Export Template", {"template_name": AQAR_EXCEL}), 1)

	def test_every_seeded_aqar_cell_is_unique_and_covers_the_answer_nodes(self):
		template = frappe.get_doc("Export Template", AQAR_EXCEL)
		cells = [(r.sheet.lower(), r.cell.upper()) for r in template.mappings]
		self.assertEqual(len(cells), len(set(cells)))
		answer_nodes = frappe.get_all(
			"Framework Node",
			{"framework": FRAMEWORK, "node_type": ["in", ["Metric", "Sub-Metric", "Extended Profile"]]},
			pluck="name",
		)
		self.assertEqual({r.node for r in template.mappings}, set(answer_nodes))
