"""M01-09 AISHE dataset: an Excel file of the AISHE sections the college holds data for, built from verified records.

Sections: institution basic details (Institution Profile), programmes with sanctioned intake and admitted students (Seat
Matrix and submitted Program Enrollments), enrolment by category, gender and state (distinct students with a submitted
Program Enrollment: the total_students data point), faculty by designation, employment type and gender (the teachers of the
full_time_teachers data point), examination results by programme (the Result Analysis report, final-year sessions: the
final_year_appeared / final_year_passed data points) and scholarships by scheme type, category and gender (the Scholarship
Beneficiaries report's rows: the scholarship_*_students data points). Faculty social category is not held on Employee, so
employment type fills the category column of the faculty table.

Where each section goes in the file is an AISHE Layout record. The official AISHE template was not available, so the seeded
layout is labelled a placeholder and the cover sheet says so. Cells are written with the Export Template helpers
(campus.iqac.export), so no cell can run as a spreadsheet formula. Once the role and company are checked, every source is read system-wide for that company (frappe.get_all), so the
figures do not depend on the exporter's other roles.
"""

import io
import re
from collections import Counter

import frappe
from frappe import _
from frappe.utils import now_datetime
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from campus.campus_finance.api import scholarship_beneficiaries
from campus.campus_platform.company_scope import permitted_companies
from campus.iqac import datapoints
from campus.iqac.access import READERS
from campus.iqac.export import COVER, parse_cell, put, sheet_name, unique_sheet_name

SECTIONS = ["Institution", "Programmes", "Enrolment", "Faculty", "Examination Results", "Scholarships"]
PLACEHOLDER_LAYOUT = "AISHE Placeholder Layout"
NOT_RECORDED = "Not recorded"
PROFILE_ITEMS = [  # label, Institution Profile field
	("AISHE Code", "aishe_code"), ("AICTE Permanent ID", "aicte_permanent_id"), ("NAAC Track ID", "naac_track_id"),
	("NIRF ID", "nirf_id"), ("Institution Status", "institution_status"), ("Affiliating University", "affiliating_university"),
	("Segment", "segment"),
]  # fmt: skip


def export_dataset(company, academic_year, layout=None):
	"""(file name, content) of the AISHE workbook of the company and year; IQAC staff and the Principal of that company."""
	frappe.only_for(READERS)
	permitted = permitted_companies(frappe.session.user)
	if permitted is not None and company not in permitted:
		frappe.throw(
			_("You may not export the data of {0}.").format(frappe.bold(company)), frappe.PermissionError
		)
	layout = frappe.get_doc("AISHE Layout", layout or PLACEHOLDER_LAYOUT)
	book = Workbook()
	cover = book.active
	cover.title = COVER
	for number, (label, value) in enumerate(cover_rows(company, academic_year, layout), 1):
		put(cover, f"A{number}", label)
		put(cover, f"B{number}", value)
	write_sections(book, layout, build(company, academic_year))
	stream = io.BytesIO()
	book.save(stream)
	return f"{re.sub(r'[^A-Za-z0-9._-]+', '_', f'AISHE-{company}-{academic_year}')}.xlsx", stream.getvalue()


def cover_rows(company, academic_year, layout):
	rows = [
		(_("Dataset"), "AISHE"), (_("Layout"), layout.name), (_("College"), company), (_("Academic Year"), academic_year),
		(_("Exported By"), frappe.session.user), (_("Exported On"), str(now_datetime().replace(microsecond=0))),
		(_("Notice"), _("Built from submitted enrolments, published results and sanctioned scholarships of the selected company.")),
	]  # fmt: skip
	if layout.is_placeholder:
		rows.append(
			(
				_("Notice"),
				_(
					"Placeholder layout: the official AISHE template was not available. Replace the layout with the official one."
				),
			)
		)
	return rows


def write_sections(book, layout, sections):
	sheets, written = {}, set()
	for row in layout.sections:
		if row.sheet.lower() not in sheets:
			sheets[row.sheet.lower()] = book.create_sheet(unique_sheet_name(book, sheet_name(row.sheet)))
		column, first_row = parse_cell(row.first_cell)
		table = sections[row.section]
		for offset, values in enumerate([table["header"], *table["rows"]]):
			for shift, value in enumerate(values):
				cell = (row.sheet.lower(), column + shift, first_row + offset)
				if cell in written:
					frappe.throw(
						_(
							"The {0} section overlaps another section on {1}; move one in the AISHE Layout."
						).format(frappe.bold(row.section), frappe.bold(row.sheet))
					)
				written.add(cell)
				put(
					sheets[row.sheet.lower()],
					f"{get_column_letter(column + shift)}{first_row + offset}",
					value,
				)


def build(company, academic_year):
	"""{section: {header, rows}} for the company and year, from submitted / published / sanctioned records."""
	return {
		section: dict(zip(("header", "rows"), builder(company, academic_year), strict=True))
		for section, builder in (
			("Institution", institution),
			("Programmes", programmes),
			("Enrolment", enrolment),
			("Faculty", faculty),
			("Examination Results", examination_results),
			("Scholarships", scholarships),
		)  # fmt: skip
	}


def institution(company, academic_year):
	profile = (
		frappe.db.get_value(
			"Institution Profile",
			{"company": company},
			[field for _label, field in PROFILE_ITEMS],
			as_dict=True,
		)
		or {}
	)
	rows = [["Institution", company], ["Academic Year", academic_year]]
	return ["Item", "Value"], rows + [[label, profile.get(field)] for label, field in PROFILE_ITEMS]


def enrollments(company, academic_year):
	return frappe.get_all(
		"Program Enrollment", {"oc_company": company, "academic_year": academic_year, "docstatus": 1},
		["student", "program", "student_category"], order_by="name",
	)  # fmt: skip


def programmes(company, academic_year):
	matrices = frappe.get_all(
		"Seat Matrix", {"company": company, "academic_year": academic_year}, ["name", "program"]
	)
	program_of = {m.name: m.program for m in matrices}
	intake = Counter()
	seats = frappe.get_all(
		"Seat Matrix Row", {"parent": ["in", list(program_of) or [""]], "quota": ["in", datapoints.counted_quotas() or [""]]},
		["parent", "sanctioned"], limit=0,
	)  # fmt: skip
	for seat in seats:
		intake[program_of[seat.parent]] += seat.sanctioned or 0
	admitted = {}
	for row in enrollments(company, academic_year):
		admitted.setdefault(row.program, set()).add(row.student)
	programs = sorted({*program_of.values(), *admitted})
	departments = dict(
		frappe.get_all("Program", {"name": ["in", programs or [""]]}, ["name", "department"], as_list=True)
	)
	rows = [[p, departments.get(p) or NOT_RECORDED, intake[p], len(admitted.get(p, ()))] for p in programs]
	return ["Programme", "Department", "Sanctioned Intake", "Admitted"], rows


def student_traits(students, company, academic_year):
	"""{student: (category, gender, state)} of `students`; the category is the one on their enrolment of the year."""
	category = {}
	for row in enrollments(company, academic_year):
		category.setdefault(row.student, row.student_category)
	people = {
		s.name: s
		for s in frappe.get_all(
			"Student", {"name": ["in", list(students) or [""]]}, ["name", "gender", "state"]
		)
	}
	return {
		s: (
			category.get(s) or NOT_RECORDED,
			people[s].gender or NOT_RECORDED,
			people[s].state or NOT_RECORDED,
		)
		for s in students
		if s in people
	}


def enrolment(company, academic_year):
	students = sorted({row.student for row in enrollments(company, academic_year)})
	counts = Counter(student_traits(students, company, academic_year).values())
	return ["Category", "Gender", "State", "Students"], [[*key, n] for key, n in sorted(counts.items())]


def faculty(company, academic_year):
	from campus.faculty import staffing

	start, end = staffing.year_bounds(academic_year)
	fields = ["designation", "employment_type", "gender", "status", "relieving_date"]
	people = [
		p
		for p in frappe.get_all("Employee", staffing.teaching_filters(company, end), fields, limit=0)
		if staffing.on_roll(p, start)
	]
	counts = Counter(
		(p.designation or NOT_RECORDED, p.employment_type or NOT_RECORDED, p.gender or NOT_RECORDED)
		for p in people
	)
	return ["Designation", "Employment Type (category)", "Gender", "Teachers"], [
		[*key, n] for key, n in sorted(counts.items())
	]


def examination_results(company, academic_year):
	from campus.examinations import coe_report_rules as rules
	from campus.examinations import coe_reports

	outcomes = coe_reports.year_outcomes(company, academic_year, final_year=True)
	tally = rules.tally(outcomes, lambda r: r["program"], lambda r: r["result"])
	rows = [
		[program or NOT_RECORDED, c["appeared"], c["passed"], rules.percent(c["passed"], c["appeared"])]
		for program, c in sorted(tally.items(), key=lambda item: item[0] or "")
	]
	if rows:
		appeared, passed = sum(r[1] for r in rows), sum(r[2] for r in rows)
		rows.append([_("Total"), appeared, passed, rules.percent(passed, appeared)])
	return ["Programme", "Appeared", "Passed", "Pass %"], rows


def scholarships(company, academic_year):
	rows = scholarship_beneficiaries(company, academic_year)
	pairs = sorted({(r.scheme_type, r.student) for r in rows})
	traits = student_traits({student for _kind, student in pairs}, company, academic_year)
	counts = Counter((kind, *traits[student][:2]) for kind, student in pairs if student in traits)
	return ["Scheme Type", "Category", "Gender", "Students"], [[*key, n] for key, n in sorted(counts.items())]
