"""M01-03 Export of an Accreditation Cycle to the Excel / PDF layout of an Export Template.

The layout is the template's mapping rows (data); this module only writes what they say. Values are the cycle's Metric
Values as the user may read them. Every cell written from user data goes through `safe_text`, so nothing in a sheet can run
as a spreadsheet formula. Excel: a Cover sheet, the mapped sheets, then a Data Template sheet per flagged computed node
listing the records behind its value (from the Metric Value snapshot) that the user may read, with a count of the others. PDF: a Part A style cover block, then the mapped
sections in the order they first appear.
"""

import html
import io
import json
import re
from collections import defaultdict

import frappe
from frappe import _
from frappe.utils import flt, now_datetime
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from campus.iqac.access import READERS

COVER = "Cover"
DATA_PREFIX = "Data "
MAX_CELL = 32000
FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")
NUMERIC = ("Number", "Percent")
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
TAGS = re.compile(r"<[^>]*>")
CELL = re.compile(r"^([A-Za-z]{1,3})([1-9][0-9]{0,6})$")
FORBIDDEN_IN_SHEET = re.compile(r"[\[\]:*?/\\]")
DATA_HEADER = ["Department", "Data Point", "Data Point Value", "Document Type", "Record"]


def safe_text(value):
	"""`value` as cell text a spreadsheet will not run as a formula: control characters dropped, length capped, and text that
	starts (after spaces) with = + - @ tab or CR gets a leading apostrophe."""
	text = CONTROL.sub("", "" if value is None else str(value))[:MAX_CELL]
	if text.lstrip(" ")[:1] in FORMULA_STARTS and text.lstrip(" "):
		return "'" + text
	return text


def plain(rich_text):
	"""A rich-text answer as plain text: markup removed, entities decoded, paragraphs separated by a space."""
	return " ".join(html.unescape(TAGS.sub(" ", rich_text or "")).split())


def sheet_name(name):
	"""`name` as a sheet name Excel accepts: forbidden characters become spaces, no leading formula sign or quote, 31 characters."""
	text = (
		FORBIDDEN_IN_SHEET.sub(" ", CONTROL.sub("", name or ""))
		.strip()
		.lstrip("=+-@")
		.strip("' ")[:31]
		.strip("' ")
	)
	return text or "Sheet"


def data_sheet_name(code):
	return sheet_name(f"{DATA_PREFIX}{code}")


def sheet_problem(name):
	"""Why `name` cannot be a mapped sheet (None when it can)."""
	if sheet_name(name) != name:
		return _("use at most 31 characters, none of [ ] : * ? / \\, and not starting with = + - @ or '.")
	if name.lower() == COVER.lower() or name.lower().startswith(DATA_PREFIX.lower()):
		return _("{0} and names starting with {1} are used by the export itself.").format(
			COVER, DATA_PREFIX.strip()
		)


def parse_cell(ref):
	"""(column, row) numbers of a cell reference like B7, None when it is not one Excel has."""
	match = CELL.match((ref or "").strip())
	if not match:
		return None
	column = 0
	for letter in match.group(1).upper():
		column = column * 26 + ord(letter) - 64
	row = int(match.group(2))
	return (column, row) if column <= 16384 and row <= 1048576 else None


def export(cycle_name, template_name):
	"""(file name, content) of the cycle exported with the template; IQAC staff and the Principal of the cycle's company."""
	frappe.only_for(READERS)
	cycle = frappe.get_doc("Accreditation Cycle", cycle_name)
	cycle.check_permission("read")
	template = frappe.get_doc("Export Template", template_name)
	template.check_permission("read")
	if template.framework != cycle.framework:
		frappe.throw(
			_("{0} is a template of {1}, not of {2}.").format(
				frappe.bold(template.name), frappe.bold(template.framework), frappe.bold(cycle.framework)
			)
		)
	exported = Exported(cycle, template)
	build = {"Excel": exported.workbook, "PDF": exported.pdf}[template.export_format]
	content = build()
	cycle.add_comment("Info", _("Exported with {0}.").format(frappe.utils.escape_html(template.name)))
	extension = "xlsx" if template.export_format == "Excel" else "pdf"
	return f"{re.sub(r'[^A-Za-z0-9._-]+', '_', f'{cycle.name}-{template.name}')}.{extension}", content


class Exported:
	"""What one export reads: the cycle, the template, the framework nodes and the Metric Values the user may read."""

	def __init__(self, cycle, template):
		self.cycle, self.template = cycle, template
		self.nodes = {
			n.name: n
			for n in frappe.get_all(
				"Framework Node",
				{"framework": cycle.framework},
				["name", "code", "title", "answer_type", "unit"],
			)
		}
		self.values = frappe.get_list(
			"Metric Value", {"cycle": cycle.name}, ["node", "department", "value", "answered", "choice", "narrative", "status", "snapshot"],
			order_by="name",
		)  # fmt: skip
		self.college = {v.node: v for v in self.values if not v.department}
		self.mappings = list(template.mappings)
		self.read_rights = {}

	def answer(self, node):
		"""The college-level answer of `node`: a number, a choice, plain narrative text, or None."""
		meta, value = self.nodes[node], self.college.get(node)
		if not value or not value.answered:
			return None
		if meta.answer_type in NUMERIC:
			return flt(value.value)
		if meta.answer_type == "Choice":
			return value.choice
		return plain(value.narrative) if meta.answer_type == "Narrative" else None

	def cover_rows(self):
		unlocked = [v for v in self.values if v.status != "Locked"]
		rows = [
			(_("Template"), self.template.name), (_("Framework"), self.cycle.framework), (_("Cycle"), self.cycle.name),
			(_("College"), self.cycle.company), (_("Academic Year"), self.cycle.academic_year), (_("Cycle Status"), _(self.cycle.status)),
			(_("Exported By"), frappe.session.user), (_("Exported On"), str(now_datetime().replace(microsecond=0))),
		]  # fmt: skip
		if self.template.is_placeholder:
			rows.append(
				(
					_("Notice"),
					_(
						"Placeholder layout: the official template was not available. Replace the mappings with the official layout."
					),
				)
			)
		if unlocked:
			rows.append(
				(
					_("Notice"),
					_("{0} of {1} values are not Locked; these figures are not final.").format(
						len(unlocked), len(self.values)
					),
				)
			)
		return rows

	# Excel

	def workbook(self):
		book = Workbook()
		cover = book.active
		cover.title = COVER
		for number, (label, value) in enumerate(self.cover_rows(), 1):
			put(cover, f"A{number}", label)
			put(cover, f"B{number}", value)
		sheets = {}
		for row in self.mappings:
			if row.sheet.lower() not in sheets:
				sheets[row.sheet.lower()] = book.create_sheet(unique_sheet_name(book, sheet_name(row.sheet)))
			sheet = sheets[row.sheet.lower()]
			put(sheet, row.cell, self.answer(row.node))
			if row.label_cell:
				put(sheet, row.label_cell, f"{self.nodes[row.node].code} {self.nodes[row.node].title}")
		for node in dict.fromkeys(row.node for row in self.mappings if row.data_sheet):
			self.data_sheet(book, node)
		stream = io.BytesIO()
		book.save(stream)
		return stream.getvalue()

	def data_sheet(self, book, node):
		"""A listing of the records behind the node's value(s); nothing for a node without a snapshot (answered by hand)."""
		rows, hidden = self.listing(node)
		if not rows and not hidden:
			return
		meta = self.nodes[node]
		sheet = book.create_sheet(unique_sheet_name(book, data_sheet_name(meta.code)))
		put(sheet, "A1", f"{meta.code} {meta.title}")
		put(
			sheet,
			"A2",
			_("Cycle {0}, status of the college value: {1}").format(
				self.cycle.name, _(self.college[node].status) if node in self.college else "-"
			),
		)
		for column, label in enumerate(DATA_HEADER, 1):
			put(sheet, f"{get_column_letter(column)}4", label)
		for number, row in enumerate(rows, 5):
			for column, value in enumerate(row, 1):
				put(sheet, f"{get_column_letter(column)}{number}", value)
		if hidden:
			put(sheet, f"A{len(rows) + 6}", _("{0} records you cannot open are not listed.").format(hidden))

	def listing(self, node):
		"""Rows of the records behind the node's values that the user may read, and how many were left out."""
		rows = []
		for value in sorted(
			(v for v in self.values if v.node == node), key=lambda v: (bool(v.department), v.department or "")
		):
			inputs = json.loads(value.snapshot or "{}").get("inputs", {})
			for key, point in inputs.items():
				for record in point.get("records") or [{}]:
					rows.append(
						[
							value.department or _("College"),
							key,
							flt(point.get("value")),
							record.get("doctype"),
							record.get("name"),
						]
					)
		names = defaultdict(set)
		for row in rows:
			if row[3]:
				names[row[3]].add(row[4])
		readable = {doctype: self.readable(doctype, wanted) for doctype, wanted in names.items()}
		shown = [row for row in rows if not row[3] or row[4] in readable[row[3]]]
		return shown, sum(len(wanted - readable[doctype]) for doctype, wanted in names.items())

	def readable(self, doctype, names):
		"""The `names` of `doctype` the user may read (one get_list per doctype per export, so the permission query applies)."""
		cached = self.read_rights.setdefault(doctype, {})
		unknown = [name for name in names if name not in cached]
		if unknown:
			allowed = set()
			if frappe.db.exists("DocType", doctype) and frappe.has_permission(doctype, "read"):
				allowed = set(
					frappe.get_list(doctype, {"name": ["in", unknown]}, pluck="name", limit_page_length=0)
				)
			cached.update({name: name in allowed for name in unknown})
		return {name for name in names if cached[name]}

	# PDF

	def pdf(self):
		from frappe.utils.pdf import get_pdf

		sections = {}
		for row in self.mappings:
			sections.setdefault(row.section.strip(), []).append(row)
		body = [
			f"<h1>{html.escape(self.template.name)}</h1>",
			f"<h2>{html.escape(self.template.cover_heading or _('Part A - Institutional Details'))}</h2>",
		]
		body.append(
			"<table>"
			+ "".join(
				f"<tr><th>{html.escape(str(k))}</th><td>{html.escape(str(v))}</td></tr>"
				for k, v in self.cover_rows()
			)
			+ "</table>"
		)
		for section, rows in sections.items():
			body.append(f"<h2 class='section'>{html.escape(section)}</h2><table>")
			for row in rows:
				meta = self.nodes[row.node]
				body.append(
					f"<tr><th>{html.escape(meta.code)}</th><td>{html.escape(meta.title)}</td><td class='v'>{html.escape(self.printed(row.node))}</td></tr>"
				)
			body.append("</table>")
		return get_pdf(
			f"<html><head><meta charset='utf-8'><style>{PDF_STYLE}</style></head><body>{''.join(body)}</body></html>"
		)

	def printed(self, node):
		answer = self.answer(node)
		if answer is None:
			return ""
		if isinstance(answer, float):
			answer = f"{answer:.2f}".rstrip("0").rstrip(".")
		unit = self.nodes[node].unit
		return f"{answer} {unit}" if unit and not isinstance(answer, str) else str(answer)


PDF_STYLE = """
body { font-family: sans-serif; font-size: 11px; } h1 { font-size: 18px; } h2 { font-size: 14px; margin-top: 18px; }
h2.section { page-break-before: always; } table { width: 100%; border-collapse: collapse; margin-top: 6px; }
th, td { border: 1px solid #999; padding: 4px 6px; text-align: left; vertical-align: top; } th { width: 12%; white-space: nowrap; }
td.v { width: 25%; white-space: pre-wrap; }
"""


def put(sheet, ref, value):
	"""Write `value` to the cell: numbers as numbers, anything else as neutralised text that is never a formula."""
	if value is None:
		return
	cell = sheet[ref]
	if isinstance(value, (int, float)) and not isinstance(value, bool):
		cell.value = value
		return
	cell.value = safe_text(value)
	cell.data_type = "s"


def unique_sheet_name(book, base):
	name, number = base, 1
	taken = {n.lower() for n in book.sheetnames}
	while name.lower() in taken:
		number += 1
		suffix = f" {number}"
		name = base[: 31 - len(suffix)] + suffix
	return name
